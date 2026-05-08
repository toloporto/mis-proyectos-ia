import io
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams
from langchain_qdrant import QdrantVectorStore
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

EMBEDDING_MODEL = "nomic-embed-text"
QDRANT_URL = "http://localhost:6333"
VECTOR_SIZE = 768
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200


def _get_embeddings():
    return OllamaEmbeddings(model=EMBEDDING_MODEL)


def _get_client():
    return QdrantClient(url=QDRANT_URL)


def _ensure_collection(client: QdrantClient, collection_name: str):
    existing = {c.name for c in client.get_collections().collections}
    if collection_name not in existing:
        client.create_collection(
            collection_name=collection_name,
            vectors_config=VectorParams(size=VECTOR_SIZE, distance=Distance.COSINE),
        )


def get_qdrant_retriever(collection_name: str):
    client = _get_client()
    store = QdrantVectorStore(
        client=client, collection_name=collection_name, embedding=_get_embeddings()
    )
    return store.as_retriever(search_kwargs={"k": 5})


def retrieve_context(query: str, collection_name: str) -> str:
    """Recupera fragmentos relevantes de Qdrant y los devuelve como texto."""
    try:
        retriever = get_qdrant_retriever(collection_name)
        docs = retriever.invoke(query)
        return "\n\n---\n\n".join(d.page_content for d in docs) if docs else ""
    except Exception as e:
        print(f"[RAG] Error recuperando contexto de '{collection_name}': {e}")
        return ""


def _extract_text(file_bytes: bytes, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "txt":
        return file_bytes.decode("utf-8", errors="ignore")
    if ext == "pdf":
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext in ("docx", "doc"):
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"Formato no soportado: .{ext}  (acepta: pdf, docx, txt)")


def ingest_document(file_bytes: bytes, filename: str, collection_name: str) -> int:
    """Indexa un archivo en Qdrant. Devuelve el número de chunks creados."""
    raw_text = _extract_text(file_bytes, filename)
    if not raw_text.strip():
        raise ValueError("El archivo no contiene texto extraíble.")

    chunks = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP
    ).split_text(raw_text)

    client = _get_client()
    _ensure_collection(client, collection_name)

    store = QdrantVectorStore(
        client=client, collection_name=collection_name, embedding=_get_embeddings()
    )
    docs = [
        Document(page_content=chunk, metadata={"source": filename, "chunk_idx": i})
        for i, chunk in enumerate(chunks)
    ]
    store.add_documents(docs)
    return len(chunks)


def delete_collection(collection_name: str):
    try:
        _get_client().delete_collection(collection_name)
    except Exception as e:
        print(f"[RAG] No se pudo eliminar '{collection_name}': {e}")
